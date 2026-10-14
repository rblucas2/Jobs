"""Exportação em Word (python-docx, a partir do modelo fixo) e PDF (LibreOffice headless)."""
from __future__ import annotations

import os
import platform
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt, RGBColor

from ..config import resolver

ESTILOS = {
    "nome": "CV Nome",
    "contactos": "CV Contactos",
    "seccao": "CV Secção",
    "entrada": "CV Entrada",
    "detalhe": "CV Detalhe",
    "ponto": "CV Ponto",
    "texto": "CV Texto",
}


def criar_modelo(destino: Path) -> Path:
    """Cria o modelo Word por defeito (estilos CV *). Pode ser editado no Word, mantendo os nomes dos estilos."""
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Pt(42)
        s.left_margin = s.right_margin = Pt(50)
    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(10.5)

    def estilo(nome, tamanho, negrito=False, cor=None, antes=0, depois=2, baseado="Normal", italico=False):
        st = doc.styles.add_style(nome, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles[baseado]
        st.font.size = Pt(tamanho)
        st.font.bold = negrito
        st.font.italic = italico
        if cor:
            st.font.color.rgb = RGBColor.from_string(cor)
        st.paragraph_format.space_before = Pt(antes)
        st.paragraph_format.space_after = Pt(depois)
        return st

    estilo(ESTILOS["nome"], 20, negrito=True, cor="1F3864", depois=0)
    estilo(ESTILOS["contactos"], 9.5, cor="595959", depois=8)
    estilo(ESTILOS["seccao"], 11.5, negrito=True, cor="1F3864", antes=10, depois=3)
    estilo(ESTILOS["entrada"], 10.5, negrito=True, antes=5, depois=0)
    estilo(ESTILOS["detalhe"], 9.5, cor="595959", italico=True, depois=2)
    estilo(ESTILOS["ponto"], 10.5, depois=1, baseado="List Bullet")
    estilo(ESTILOS["texto"], 10.5, depois=4)
    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destino)
    return destino


def _limpar_corpo(doc) -> None:
    corpo = doc.element.body
    for el in list(corpo):
        if not el.tag.endswith("sectPr"):
            corpo.remove(el)


def gerar_docx(blocos: list[tuple[str, str]], modelo: Path, destino: Path) -> Path:
    if not modelo.exists():
        raise FileNotFoundError(f"Modelo de CV não encontrado: {modelo}")
    doc = Document(str(modelo))
    _limpar_corpo(doc)
    existentes = {s.name for s in doc.styles}
    for estilo, texto in blocos:
        nome = ESTILOS.get(estilo, "Normal")
        if nome not in existentes:
            nome = "List Bullet" if estilo == "ponto" else "Normal"
        doc.add_paragraph(texto, style=nome)
    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(destino))
    return destino


def encontrar_soffice(cfg: dict) -> str | None:
    if cfg["cv"].get("soffice"):
        return cfg["cv"]["soffice"]
    for nome in ("soffice", "libreoffice"):
        if shutil.which(nome):
            return shutil.which(nome)
    candidatos = {
        "Windows": [r"C:\Program Files\LibreOffice\program\soffice.exe",
                    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"],
        "Darwin": ["/Applications/LibreOffice.app/Contents/MacOS/soffice"],
    }.get(platform.system(), [])
    return next((c for c in candidatos if os.path.exists(c)), None)


def converter_pdf(docx: Path, cfg: dict) -> Path | None:
    """Converte com o LibreOffice em modo headless. Devolve None se não estiver instalado."""
    soffice = encontrar_soffice(cfg)
    if not soffice:
        return None
    subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
        check=True, capture_output=True, timeout=180,
    )
    pdf = docx.with_suffix(".pdf")
    return pdf if pdf.exists() else None


def caminho_modelo(cfg: dict) -> Path:
    return resolver(cfg, cfg["cv"]["modelo_docx"])
