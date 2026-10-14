"""Extração do texto do CV atual (PDF ou Word) e proposta de perfil-base estruturado.

A proposta é sempre revista e validada pelo utilizador. Para reduzir o risco de erros
(incluindo do modelo local, se for usado), cada item proposto é comparado com o texto
do CV original e os que não forem encontrados são assinalados para revisão.
"""
from __future__ import annotations

import io
import re
from difflib import SequenceMatcher
from pathlib import Path

from .. import llm
from ..texto import normalizar
from .modelo import (Certificacao, Competencia, DadosPessoais, Experiencia, Formacao, ItemTexto, Lingua,
                     Perfil, indexar)
from .versoes import atribuir_ids

# ---------------------------------------------------------------- leitura de ficheiros


def extrair_texto(nome: str, dados: bytes) -> str:
    ext = Path(nome).suffix.lower()
    if ext == ".pdf":
        import pdfplumber
        with pdfplumber.open(io.BytesIO(dados)) as pdf:
            return "\n".join((p.extract_text() or "") for p in pdf.pages)
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(dados))
        linhas = []
        for p in d.paragraphs:
            t = p.text.strip()
            if t:
                estilo = (p.style.name or "").lower() if p.style is not None else ""
                linhas.append(("• " + t) if "list" in estilo else t)
        for tabela in d.tables:
            for linha in tabela.rows:
                for celula in linha.cells:
                    for p in celula.paragraphs:
                        if p.text.strip():
                            linhas.append(p.text.strip())
        return "\n".join(linhas)
    if ext in (".txt", ".md"):
        return dados.decode("utf-8", errors="replace")
    raise ValueError("Formato não suportado: use PDF ou Word (.docx).")


# ---------------------------------------------------------------- heurística por secções

SECCOES = {
    "resumo": r"(perfil|resumo|sobre mim|summary|profile|about me|objetivo|objective)",
    "experiencia": r"(experi[eê]ncia( profissional)?|percurso profissional|(professional |work )?experience|employment)",
    "formacao": r"(forma[cç][aã]o( acad[eé]mica)?|educa[cç][aã]o|education|habilita[cç][oõ]es)",
    "competencias": r"(compet[eê]ncias|skills|conhecimentos|ferramentas|inform[aá]tica|technical skills)",
    "linguas": r"(l[ií]nguas|idiomas|languages)",
    "certificacoes": r"(certifica[cç][oõ]es|certifica[cç][aã]o|certifications|cursos|courses)",
}
CABECALHO = {k: re.compile(rf"^\s*{v}\s*:?\s*$", re.I) for k, v in SECCOES.items()}

MESES = {
    "jan": 1, "fev": 2, "feb": 2, "mar": 3, "abr": 4, "apr": 4, "mai": 5, "may": 5, "jun": 6,
    "jul": 7, "ago": 8, "aug": 8, "set": 9, "sep": 9, "out": 10, "oct": 10, "nov": 11, "dez": 12, "dec": 12,
}
_MES = r"(?:(?:jan|fev|feb|mar|abr|apr|mai|may|jun|jul|ago|aug|set|sep|out|oct|nov|dez|dec)[a-zç]*\.?\s*|\d{1,2}[/.-])?"
INTERVALO = re.compile(
    rf"({_MES}\d{{4}})\s*(?:-|–|—|a|até|to)\s*({_MES}\d{{4}}|atual|atualmente|presente|present|current|hoje|now)",
    re.I,
)
MARCADOR = re.compile(r"^\s*[•▪●◦\-–*·]\s*")
RESULTADO = re.compile(r"(\d|%|aument|reduz|poupan|cresc|melhor|increas|reduc|sav|grew|improv)", re.I)
NIVEL = re.compile(r"\b(A1|A2|B1|B2|C1|C2|nativ[oa]|native|materna|fluente|fluent|avançado|advanced|"
                   r"intermédio|intermediate|básico|basic)\b", re.I)


def _data(token: str) -> str | None:
    t = token.strip().lower()
    if re.match(r"^(atual|atualmente|presente|present|current|hoje|now)$", t):
        return "atual"
    ano = re.search(r"(\d{4})", t)
    if not ano:
        return None
    mes = None
    m = re.match(r"^([a-zç]{3})", t)
    if m and m.group(1) in MESES:
        mes = MESES[m.group(1)]
    m2 = re.match(r"^(\d{1,2})[/.-]", t)
    if m2 and 1 <= int(m2.group(1)) <= 12:
        mes = int(m2.group(1))
    return f"{ano.group(1)}-{mes:02d}" if mes else ano.group(1)


def _dividir_seccoes(texto: str) -> dict[str, list[str]]:
    seccoes: dict[str, list[str]] = {"_topo": []}
    atual = "_topo"
    for linha in texto.splitlines():
        l = linha.strip()
        if not l:
            continue
        nova = next((k for k, rx in CABECALHO.items() if rx.match(l)), None)
        if nova:
            atual = nova
            seccoes.setdefault(atual, [])
            continue
        seccoes.setdefault(atual, []).append(l)
    return seccoes


def _cargo_empresa(linhas: list[str]) -> tuple[str, str, str]:
    """De 1–3 linhas de cabeçalho, obtém (cargo, empresa, local)."""
    linhas = [INTERVALO.sub("", l).strip(" |,–-·") for l in linhas]
    linhas = [l for l in linhas if l]
    if not linhas:
        return "", "", ""
    if len(linhas) == 1:
        partes = [p.strip() for p in re.split(r"\s+(?:\||–|—|-|@|·)\s+|,\s+", linhas[0]) if p.strip()]
        partes += ["", "", ""]
        return partes[0], partes[1], partes[2]
    cargo, empresa = linhas[0], linhas[1]
    local = linhas[2] if len(linhas) > 2 else ""
    if "," in empresa and not local:
        empresa, local = [x.strip() for x in empresa.split(",", 1)]
    return cargo, empresa, local


def _entradas(linhas: list[str]) -> list[tuple[list[str], str | None, str | None, list[str]]]:
    """Agrupa linhas em entradas (cabeçalho, início, fim, pontos), cada uma marcada por um intervalo de datas.

    As linhas sem marcador imediatamente antes de uma linha com datas (até 2) são o
    cabeçalho dessa entrada; as restantes são pontos da entrada anterior.
    """
    entradas: list[tuple[list[str], str | None, str | None, list[str]]] = []
    pendentes: list[str] = []
    atual = None

    def descarregar(n_manter: int) -> list[str]:
        resto = pendentes[:len(pendentes) - n_manter] if n_manter else list(pendentes)
        if atual is not None:
            atual[3].extend(resto)
        manter = pendentes[len(pendentes) - n_manter:] if n_manter else []
        pendentes.clear()
        return manter

    for l in linhas:
        m = INTERVALO.search(l)
        if m:
            cab = descarregar(min(2, len(pendentes)))
            atual = (cab + [l], _data(m.group(1)), _data(m.group(2)), [])
            entradas.append(atual)
        elif MARCADOR.match(l):
            descarregar(0)
            if atual is not None:
                atual[3].append(MARCADOR.sub("", l))
        elif atual is not None and not atual[3] and not pendentes and len(atual[0]) < 3:
            atual[0].append(l)  # linha de cabeçalho a seguir às datas (ex.: local)
        elif pendentes == [] and atual is not None and atual[3] and l[:1].islower():
            atual[3][-1] += " " + l  # continuação da linha anterior
        else:
            pendentes.append(l)
    descarregar(0)
    return entradas


def _lista(linhas: list[str]) -> list[str]:
    itens = []
    for l in linhas:
        for p in re.split(r"[,;•▪●|]|\s{2,}", MARCADOR.sub("", l)):
            p = p.strip(" .:")
            if p and len(p) <= 80:
                itens.append(p)
    return itens


def proposta_por_regras(texto: str) -> Perfil:
    s = _dividir_seccoes(texto)
    topo = s.get("_topo", [])
    tudo = "\n".join(topo)
    dp = DadosPessoais(
        nome=topo[0] if topo else "",
        email=(re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", texto) or [""])[0],
        telefone=(re.search(r"(\+?\d[\d ]{8,15}\d)", tudo) or [""])[0].strip(),
        linkedin=(re.search(r"(https?://)?(www\.)?linkedin\.com/in/[\w-]+/?", texto) or [""])[0],
    )
    perfil = Perfil(dados_pessoais=dp)
    perfil.resumo = [ItemTexto(texto=l, tipo="resumo") for l in s.get("resumo", [])]

    for cab, ini, fim, pontos in _entradas(s.get("experiencia", [])):
        cargo, empresa, local = _cargo_empresa(cab)
        perfil.experiencia.append(Experiencia(
            cargo=cargo or "?", empresa=empresa or "?", local=local, data_inicio=ini, data_fim=fim,
            itens=[ItemTexto(texto=p, tipo="resultado" if RESULTADO.search(p) else "responsabilidade") for p in pontos],
        ))
    for cab, ini, fim, pontos in _entradas(s.get("formacao", [])):
        grau, inst, _ = _cargo_empresa(cab)
        perfil.formacao.append(Formacao(
            grau=grau or "?", instituicao=inst, data_inicio=ini, data_fim=fim,
            itens=[ItemTexto(texto=p, tipo="detalhe") for p in pontos],
        ))
    for nome in _lista(s.get("competencias", [])):
        perfil.competencias.append(Competencia(nome=nome))
    for l in s.get("linguas", []):
        for parte in re.split(r"[,;•|]", MARCADOR.sub("", l)):
            parte = parte.strip()
            if not parte:
                continue
            nivel = NIVEL.search(parte)
            nome = re.split(r"\s*[-–:(]\s*|\s+(?=" + (re.escape(nivel.group(0)) if nivel else "$^") + ")", parte)[0].strip()
            perfil.linguas.append(Lingua(nome=nome, nivel=nivel.group(0) if nivel else ""))
    for l in s.get("certificacoes", []):
        l = MARCADOR.sub("", l).strip()
        ano = re.search(r"\b(19|20)\d{2}\b", l)
        perfil.certificacoes.append(Certificacao(nome=re.sub(r"\(?\b(19|20)\d{2}\b\)?", "", l).strip(" -–,"),
                                                 data=ano.group(0) if ano else None))
    return atribuir_ids(perfil)


# ---------------------------------------------------------------- com modelo local (opcional)

SISTEMA_EXTRACAO = (
    "Estruturas o texto de um CV em JSON. Copia o texto tal como está no CV: não reformules, "
    "não traduzas, não resumas e não acrescentes nada que não esteja escrito. Datas no formato "
    "AAAA-MM ou AAAA; usa 'atual' para cargos em curso. Deixa os campos 'id' vazios."
)


def proposta_por_llm(cfg: dict, texto: str) -> Perfil:
    p = llm.gerar_json(cfg, SISTEMA_EXTRACAO, f"Texto do CV:\n\n{texto}", Perfil)
    for i in p.resumo + [i for e in p.experiencia for i in e.itens] + [i for f in p.formacao for i in f.itens]:
        i.id = ""
    for grupo in (p.experiencia, p.formacao, p.competencias, p.linguas, p.certificacoes):
        for x in grupo:
            x.id = ""
    return atribuir_ids(p)


# ---------------------------------------------------------------- verificação contra o CV original

def itens_nao_encontrados(perfil: Perfil, texto_cv: str, limiar: float = 0.85) -> list[tuple[str, str]]:
    """Itens de texto da proposta que não aparecem (quase) literalmente no CV original."""
    alvo = normalizar(texto_cv)
    falhas = []
    for item in indexar(perfil).values():
        if item.tipo in ("experiencia", "formacao"):
            continue
        t = normalizar(item.texto)
        if not t or t in alvo:
            continue
        sm = SequenceMatcher(None, t, alvo, autojunk=False)
        bloco = sm.find_longest_match(0, len(t), 0, len(alvo))
        if bloco.size / max(1, len(t)) < limiar:
            falhas.append((item.id, item.texto))
    return falhas


def propor_perfil(cfg: dict, texto: str, usar_llm: bool | None = None) -> tuple[Perfil, str, list[tuple[str, str]]]:
    """Devolve (proposta, método, itens a rever)."""
    metodo = "regras"
    if usar_llm is None:
        usar_llm = llm.disponivel(cfg)
    perfil = None
    if usar_llm:
        try:
            perfil, metodo = proposta_por_llm(cfg, texto), "llm"
        except (llm.LLMIndisponivel, ValueError):
            perfil = None
    if perfil is None:
        perfil = proposta_por_regras(texto)
    return perfil, metodo, itens_nao_encontrados(perfil, texto)
