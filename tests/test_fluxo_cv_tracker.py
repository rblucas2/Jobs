"""M4 (fluxo: aprovação humana e exportação) e M5 (tracker e métricas)."""
from __future__ import annotations

import csv
import io
from datetime import date
from pathlib import Path

import docx
import pytest

from src import db
from src.cv import fluxo, gerador
from src.cv.exportacao import encontrar_soffice
from src.cv.schemas import CVAdaptado
from src.cv.validador import pode_aprovar
from src.tracker import estados, metricas


@pytest.fixture
def vaga_id(conn):
    return db.inserir_vaga(conn, {
        "cargo": "Analista de Controlo de Gestão", "empresa": "Retalho Alfa", "localizacao": "Lisboa",
        "chave_normalizada": "retalho alfa|analista|lisboa", "fonte": "LinkedIn",
        "descricao": "Reporting mensal em Excel e Power BI, automatização em VBA.", "descricao_completa": True,
    })


def test_fallback_por_regras_gera_cv_totalmente_verificado(cfg, perfil, vaga_id, conn):
    cv = gerador.gerar_cv(perfil, db.obter_vaga(conn, vaga_id), cfg, usar_llm=False)
    assert cv.metodo == "regras" and pode_aprovar(cv)
    assert cv.experiencia[0].experiencia_id == "exp_intermarche"  # mais recente primeiro
    assert cv.competencias_ids[:3] == ["comp_excel", "comp_power_bi", "comp_vba"]  # relevantes primeiro


def test_sem_perfil_validado_nao_gera(cfg, conn, vaga_id):
    from src.perfil.versoes import PerfilNaoValidado
    with pytest.raises(PerfilNaoValidado):
        fluxo.criar_rascunho(conn, cfg, vaga_id, usar_llm=False)


def test_aprovacao_exige_confirmacao_e_pontos_verificados(cfg_com_perfil, conn, vaga_id):
    cfg = cfg_com_perfil
    cv_id = fluxo.criar_rascunho(conn, cfg, vaga_id, usar_llm=False)
    reg = db.obter_cv(conn, cv_id)

    # Edição com facto inventado: revalidada e bloqueada.
    conteudo = reg["conteudo"]
    conteudo["experiencia"][0]["pontos"][0]["texto"] = "Gestão de uma equipa de 12 analistas."
    doc = fluxo.editar(conn, cfg, cv_id, conteudo)
    assert doc.experiencia[0].pontos[0].estado == "nao_verificado"

    with pytest.raises(fluxo.AprovacaoRecusada):
        fluxo.aprovar(conn, cfg, cv_id, confirmado=True)
    with pytest.raises(fluxo.AprovacaoRecusada):
        fluxo.exportar(conn, cfg, cv_id)  # rascunho não se exporta

    # Remover o ponto inválido permite aprovar, mas só com confirmação explícita.
    conteudo["experiencia"][0]["pontos"].pop(0)
    fluxo.editar(conn, cfg, cv_id, conteudo)
    with pytest.raises(fluxo.AprovacaoRecusada):
        fluxo.aprovar(conn, cfg, cv_id, confirmado=False)
    fluxo.aprovar(conn, cfg, cv_id, confirmado=True)
    assert db.obter_cv(conn, cv_id)["estado"] == "aprovado"

    # Um CV aprovado é imutável.
    with pytest.raises(PermissionError):
        fluxo.editar(conn, cfg, cv_id, conteudo)


def test_exportacao_exige_modelo_aprovado_e_liga_ao_tracker(cfg_com_perfil, conn, vaga_id):
    cfg = cfg_com_perfil
    cv_id = fluxo.criar_rascunho(conn, cfg, vaga_id, usar_llm=False)
    fluxo.aprovar(conn, cfg, cv_id, confirmado=True)
    with pytest.raises(fluxo.AprovacaoRecusada):
        fluxo.exportar(conn, cfg, cv_id)
    cfg["cv"]["modelo_aprovado"] = True
    tem_libreoffice = encontrar_soffice(cfg) is not None
    caminho_docx, caminho_pdf = fluxo.exportar(conn, cfg, cv_id)
    texto = "\n".join(p.text for p in docx.Document(caminho_docx).paragraphs)
    assert "Pessoa Exemplo" in texto and "Analista de Controlo de Gestão — Intermarché" in texto
    assert "set. 2021 – Atual" in texto
    estilos = {p.style.name for p in docx.Document(caminho_docx).paragraphs}
    assert {"CV Nome", "CV Secção", "CV Ponto"} <= estilos
    if tem_libreoffice:
        assert caminho_pdf and Path(caminho_pdf).exists() and Path(caminho_pdf).stat().st_size > 1000
    else:
        assert caminho_pdf is None  # sem LibreOffice o Word continua a ser exportado
    assert db.obter_candidatura(conn, vaga_id)["cv_id"] == cv_id


def test_carta_por_regras(cfg_com_perfil, conn, vaga_id):
    cid = fluxo.criar_rascunho(conn, cfg_com_perfil, vaga_id, tipo="carta", usar_llm=False)
    reg = db.obter_cv(conn, cid)
    assert reg["tipo"] == "carta" and reg["conteudo"]["paragrafos"]
    fluxo.aprovar(conn, cfg_com_perfil, cid, confirmado=True)
    bl = fluxo.blocos(conn, cfg_com_perfil, db.obter_cv(conn, cid))
    assert bl[-2] == ("texto", "Com os melhores cumprimentos,")


# ---------------------------------------------------------------- tracker

def test_transicoes_e_followup(cfg, conn, vaga_id):
    estados.mudar_estado(conn, cfg, vaga_id, estados.CANDIDATADO, data=date(2026, 10, 1))
    c = db.obter_candidatura(conn, vaga_id)
    assert c["data_candidatura"] == "2026-10-01" and c["data_followup"] == "2026-10-11"
    with pytest.raises(estados.TransicaoInvalida):
        estados.mudar_estado(conn, cfg, vaga_id, estados.POR_CANDIDATAR)
    assert estados.followups_em_atraso(conn, hoje=date(2026, 10, 14))[0]["vaga_id"] == vaga_id
    estados.mudar_estado(conn, cfg, vaga_id, estados.ENTREVISTA, data=date(2026, 10, 12))
    estados.mudar_estado(conn, cfg, vaga_id, estados.REJEITADO)
    c = db.obter_candidatura(conn, vaga_id)
    assert c["data_entrevista"] == "2026-10-12" and c["data_followup"] is None


def test_dias_followup_configuravel(cfg, conn, vaga_id):
    cfg["tracker"]["dias_followup"] = 5
    estados.mudar_estado(conn, cfg, vaga_id, estados.CANDIDATADO, data=date(2026, 10, 1))
    assert db.obter_candidatura(conn, vaga_id)["data_followup"] == "2026-10-06"


def test_so_cv_aprovado_pode_ser_marcado_como_usado(cfg_com_perfil, conn, vaga_id):
    cv_id = fluxo.criar_rascunho(conn, cfg_com_perfil, vaga_id, usar_llm=False)
    with pytest.raises(estados.TransicaoInvalida):
        estados.mudar_estado(conn, cfg_com_perfil, vaga_id, estados.CANDIDATADO, cv_id=cv_id)
    fluxo.aprovar(conn, cfg_com_perfil, cv_id, confirmado=True)
    estados.mudar_estado(conn, cfg_com_perfil, vaga_id, estados.CANDIDATADO, cv_id=cv_id)
    assert db.obter_candidatura(conn, vaga_id)["cv_id"] == cv_id


def test_metricas_e_csv(cfg, conn, vaga_id):
    outra = db.inserir_vaga(conn, {"cargo": "Valuation Analyst", "empresa": "B", "localizacao": "Porto",
                                   "chave_normalizada": "b", "fonte": "Indeed"})
    estados.mudar_estado(conn, cfg, vaga_id, estados.CANDIDATADO)
    estados.mudar_estado(conn, cfg, outra, estados.CANDIDATADO)
    estados.mudar_estado(conn, cfg, outra, estados.ENTREVISTA)
    m = metricas.metricas(conn, cfg)
    assert m["candidaturas"] == 2 and m["respondidas"] == 1 and m["taxa_resposta"] == 0.5
    assert m["entrevistas_por_fonte"] == {"Indeed": 1}
    assert m["entrevistas_por_tipo"] == {"Valuation": 1}
    dados = metricas.exportar_csv(conn, cfg)
    assert dados.startswith("﻿".encode("utf-8"))
    linhas = list(csv.DictReader(io.StringIO(dados.decode("utf-8-sig")), delimiter=";"))
    assert {l["estado"] for l in linhas} == {"Candidatado", "Entrevista"}
    assert {l["tipo_vaga"] for l in linhas} == {"Outra", "Valuation"}


def test_cv_schema_ida_e_volta(perfil, cfg, conn, vaga_id):
    cv = gerador.gerar_cv(perfil, db.obter_vaga(conn, vaga_id), cfg, usar_llm=False)
    assert CVAdaptado.model_validate(cv.model_dump()) == cv


def test_saida_do_modelo_local_com_factos_inventados_e_bloqueada(perfil, cfg, conn, vaga_id, monkeypatch):
    """Caminho com Ollama (resposta simulada): o validador corre sobre a saída do modelo."""
    from src import llm
    from src.cv.schemas import CVLLM

    resposta = CVLLM.model_validate({
        "resumo": [{"texto": "Analista com 12 anos de experiência em M&A.", "source_ids": ["resumo_01"]}],
        "experiencia": [{"experiencia_id": "exp_intermarche", "pontos": [
            {"texto": "Reporting mensal de vendas e margens para 45 lojas, em Excel e Power BI.",
             "source_ids": ["exp_intermarche_01"]},
            {"texto": "Liderança de uma equipa de 6 analistas na Sonae.", "source_ids": ["exp_intermarche_99"]},
        ]}],
        "formacao_ids": ["form_nova_sbe"], "competencias_ids": ["comp_excel"],
        "linguas_ids": ["ling_ingles"], "certificacoes_ids": ["cert_cfa"],
    })
    monkeypatch.setattr(llm, "gerar_json", lambda *a, **k: resposta)
    cv = gerador.gerar_cv(perfil, db.obter_vaga(conn, vaga_id), cfg, usar_llm=True)
    assert cv.metodo == "llm"
    assert cv.resumo[0].estado == "nao_verificado"
    assert [p.estado for p in cv.experiencia[0].pontos] == ["verificado", "nao_verificado"]
    assert any("cert_cfa" in e for e in cv.erros)
    assert not pode_aprovar(cv)


def test_ollama_indisponivel_recorre_as_regras(perfil, cfg, conn, vaga_id, monkeypatch):
    from src import llm

    def falha(*a, **k):
        raise llm.LLMIndisponivel("sem servidor")
    monkeypatch.setattr(llm, "gerar_json", falha)
    cv = gerador.gerar_cv(perfil, db.obter_vaga(conn, vaga_id), cfg, usar_llm=True)
    assert cv.metodo == "regras" and pode_aprovar(cv)
