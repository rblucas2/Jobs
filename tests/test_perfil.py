"""M2 — perfil-base: IDs estáveis, versões, validação e extração do CV."""
from __future__ import annotations

import io

import docx
import pytest
from pydantic import ValidationError

from src.perfil import versoes
from src.perfil.extracao import extrair_texto, itens_nao_encontrados, proposta_por_regras
from src.perfil.modelo import Perfil, anos_experiencia, todos_ids

CV_TEXTO = """Pessoa Exemplo
Lisboa | exemplo@example.com | +351 900 000 000

Perfil
Analista financeiro com experiência em controlo de gestão.

Experiência Profissional
Analista de Controlo de Gestão | Intermarché
Set 2021 – Atual
• Elaboração do reporting mensal para 45 lojas.
• Redução de 15% no tempo de fecho mensal.
Estagiário de Auditoria | Auditora Exemplo
Fev 2020 - Jul 2021
• Apoio a trabalhos de auditoria de 8 clientes.

Formação
Mestrado em Finanças, Universidade Exemplo
2018 - 2020

Competências
Excel, Power BI, VBA

Línguas
Português – nativo; Inglês – C1

Certificações
CFA Level I (2022)
"""


def test_ids_repetidos_sao_rejeitados(perfil):
    dados = perfil.model_dump()
    dados["competencias"][1]["id"] = "comp_excel"
    with pytest.raises(ValidationError):
        Perfil.model_validate(dados)


def test_atribuir_ids_nao_altera_ids_existentes(perfil):
    antes = todos_ids(perfil)
    perfil.experiencia[0].itens.append(perfil.experiencia[0].itens[0].model_copy(update={"id": "", "texto": "Novo"}))
    versoes.atribuir_ids(perfil)
    depois = todos_ids(perfil)
    assert [i for i in depois if i in antes] == antes
    assert "exp_intermarche_04" in depois


def test_versoes_validada_e_imutavel(cfg, perfil):
    n = versoes.guardar_nova_versao(cfg, perfil)
    with pytest.raises(versoes.PerfilNaoValidado):
        versoes.perfil_validado(cfg)
    versoes.validar_versao(cfg, n)
    assert versoes.perfil_validado(cfg).versao == n
    with pytest.raises(PermissionError):
        versoes.atualizar_rascunho(cfg, n, perfil)
    n2 = versoes.guardar_nova_versao(cfg, perfil, "alteração")
    assert n2 == n + 1 and versoes.versao_validada(cfg) == n  # a nova versão não fica validada sozinha


def test_yaml_ida_e_volta(perfil):
    p2 = versoes.de_yaml(versoes.para_yaml(perfil), 1)
    assert p2.model_dump() == perfil.model_dump()


def test_data_invalida_rejeitada(perfil):
    dados = perfil.model_dump()
    dados["experiencia"][0]["data_inicio"] = "09/2021"
    with pytest.raises(ValidationError):
        Perfil.model_validate(dados)


def test_anos_experiencia_sem_sobreposicao(perfil):
    dados = perfil.model_dump()
    dados["experiencia"][0]["data_fim"] = "2023-08"  # 2021-09..2023-08 = 2 anos
    dados["experiencia"][1].update(data_inicio="2020-09", data_fim="2021-08")  # 1 ano
    assert anos_experiencia(Perfil.model_validate(dados)) == 3.0


def test_proposta_por_regras():
    p = proposta_por_regras(CV_TEXTO)
    assert p.dados_pessoais.nome == "Pessoa Exemplo"
    assert p.dados_pessoais.email == "exemplo@example.com"
    assert [(e.cargo, e.empresa, e.data_inicio, e.data_fim) for e in p.experiencia] == [
        ("Analista de Controlo de Gestão", "Intermarché", "2021-09", "atual"),
        ("Estagiário de Auditoria", "Auditora Exemplo", "2020-02", "2021-07"),
    ]
    assert [i.texto for i in p.experiencia[0].itens] == [
        "Elaboração do reporting mensal para 45 lojas.", "Redução de 15% no tempo de fecho mensal."]
    assert p.experiencia[0].itens[1].tipo == "resultado"
    assert p.experiencia[0].id == "exp_intermarche" and p.experiencia[0].itens[0].id == "exp_intermarche_01"
    assert p.formacao[0].grau == "Mestrado em Finanças"
    assert [c.nome for c in p.competencias] == ["Excel", "Power BI", "VBA"]
    assert [(l.nome, l.nivel) for l in p.linguas] == [("Português", "nativo"), ("Inglês", "C1")]
    assert p.certificacoes[0].nome == "CFA Level I" and p.certificacoes[0].data == "2022"
    assert itens_nao_encontrados(p, CV_TEXTO) == []


def test_itens_alterados_sao_assinalados():
    p = proposta_por_regras(CV_TEXTO)
    p.experiencia[0].itens[0].texto = "Gestão de uma equipa de 20 pessoas."
    falhas = itens_nao_encontrados(p, CV_TEXTO)
    assert [f[0] for f in falhas] == ["exp_intermarche_01"]


def test_extrair_texto_docx():
    d = docx.Document()
    d.add_paragraph("Pessoa Exemplo")
    d.add_paragraph("Reporting mensal", style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    t = extrair_texto("cv.docx", buf.getvalue())
    assert "Pessoa Exemplo" in t and "• Reporting mensal" in t
